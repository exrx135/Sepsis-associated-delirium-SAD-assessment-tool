# Local data and generated outputs stay outside the source repository.
from pathlib import Path
import os
_SAD_REPO = Path(__file__).resolve().parents[2]
_SAD_PRIVATE = Path(os.environ.get("SAD_PRIVATE_ROOT", str(_SAD_REPO.parent / "study_code_and_data" / "final_analysis"))).expanduser().resolve()
def _sad_private_path(*parts):
    return _SAD_PRIVATE.joinpath(*parts)

from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = _sad_private_path("publication_tables")
OUT.mkdir(parents=True, exist_ok=True)
data = json.loads((OUT / 'table_2_transcript.json').read_text(encoding='utf-8'))
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.65)
for name in ['Normal', 'Title']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.bold = True
p = doc.add_paragraph(data['title'], style='Title')
p.paragraph_format.space_after = Pt(10)
table = doc.add_table(rows=1, cols=5)
table.autofit = False
widths = [2.35, 1.3, 0.85, 0.9, 1.55]
headers = ['Recorded clinical support\nat sepsis diagnosis', 'Predicted\nrisk group', 'Patients, n', 'SAD\nevents, n', 'Observed SAD\nincidence, %']
for c, text, width in zip(table.rows[0].cells, headers, widths):
    c.width = Inches(width)
    c.text = text
for row in data['rows']:
    cells = table.add_row().cells
    for c, key, width in zip(cells, row.keys(), widths):
        c.width = Inches(width)
        c.text = str(row[key])
for ri, row in enumerate(table.rows):
    for ci, cell in enumerate(row.cells):
        cell.vertical_alignment = 1
        pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            item = OxmlElement('w:' + edge)
            for attr, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                item.set(qn('w:' + attr), val)
            borders.append(item)
        pr.append(borders)
        margins = OxmlElement('w:tcMar')
        for edge in ['top', 'left', 'bottom', 'right']:
            item = OxmlElement('w:' + edge)
            item.set(qn('w:w'), '85')
            item.set(qn('w:type'), 'dxa')
            margins.append(item)
        pr.append(margins)
        if ri == 0:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8E8E8')
            pr.append(shade)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = 0 if ci == 0 else 1
            for run in paragraph.runs:
                run.bold = ri == 0
                run.font.size = Pt(10)
    trpr = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    trpr.append(no_split)
repeat = OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(repeat)
footnote = doc.add_paragraph(data['footnote'].replace('Values correspond to Figure 6.', 'Values correspond to Figure 5.', 1))
footnote.paragraph_format.space_before = Pt(9)
for run in footnote.runs:
    run.font.size = Pt(9.5)
doc.save(OUT / 'Table_2.docx')
print(OUT / 'Table_2.docx')
